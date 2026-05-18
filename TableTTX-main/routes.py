from flask import render_template, request, jsonify, redirect, url_for, flash, session, send_file
from app import app
from models import data_manager
from solver import timetable_solver
from ai_assistant import ai_assistant
from datetime import datetime
import logging
from io import BytesIO
from functools import wraps
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# In-memory timetable cache (Flask cookie sessions can't hold large timetable data)
_timetable_cache = {}

def _get_current_timetable():
    """Get current timetable from cache, session, or most recent saved timetable."""
    # 1. In-memory cache (fastest)
    if _timetable_cache.get('timetable'):
        return _timetable_cache['timetable']
    # 2. Session (may work for small timetables)
    tt = session.get('current_timetable')
    if tt:
        _timetable_cache['timetable'] = tt
        return tt
    # 3. Fall back to most recently saved timetable from disk
    saved = data_manager.get_saved_timetables()
    if saved:
        tt = saved[0].get('timetable', {})
        if tt:
            _timetable_cache['timetable'] = tt
            return tt
    return {}

# Authentication credentials
ADMIN_USERNAME = 'admin'
ADMIN_PASSWORD = 'admin@123'

def login_required(f):
    """Decorator to require login for routes"""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'logged_in' not in session or not session['logged_in']:
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated_function

@app.route('/login', methods=['GET', 'POST'])
def login():
    """Login page and authentication"""
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        
        if username == ADMIN_USERNAME and password == ADMIN_PASSWORD:
            session['logged_in'] = True
            session['username'] = username
            flash('Login successful!', 'success')
            return redirect(url_for('index'))
        else:
            flash('Invalid username or password. Please try again.', 'error')
    
    # If already logged in, redirect to home
    if session.get('logged_in'):
        return redirect(url_for('index'))
    
    return render_template('login.html')

@app.route('/logout')
def logout():
    """Logout and clear session"""
    _timetable_cache.clear()
    session.clear()
    flash('You have been logged out successfully.', 'info')
    return redirect(url_for('login'))

@app.route('/')
@login_required
def index():
    """Home page"""
    return render_template('index.html')

@app.route('/configure')
@login_required
def configure():
    """Academic structure configuration page"""
    structure = data_manager.get_academic_structure()
    return render_template('configure.html', structure=structure)

@app.route('/configure', methods=['POST'])
@login_required
def update_configure():
    """Update academic structure configuration"""
    try:
        structure = {
            "years": int(request.form.get('years', 3)),
            "divisions_per_year": int(request.form.get('divisions_per_year', 2)),
            "batches_per_division": int(request.form.get('batches_per_division', 3)),
            "theory_duration": int(request.form.get('theory_duration', 60)),
            "practical_duration": int(request.form.get('practical_duration', 120)),
            "time_slots": request.form.get('time_slots', '').split('\n'),
            "days": request.form.getlist('days')
        }
        
        # Clean up time slots (remove empty lines)
        structure["time_slots"] = [slot.strip() for slot in structure["time_slots"] if slot.strip()]
        
        if not structure["time_slots"]:
            structure["time_slots"] = [
                "09:00-10:00", "10:00-11:00", "11:15-12:15", "12:15-13:15",
                "14:15-15:15", "15:15-16:15", "16:30-17:30", "17:30-18:30"
            ]
        
        if not structure["days"]:
            structure["days"] = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"]
        
        data_manager.update_academic_structure(structure)
        flash('Academic structure updated successfully!', 'success')
    except Exception as e:
        flash(f'Error updating configuration: {str(e)}', 'error')
    
    return redirect(url_for('configure'))

@app.route('/subjects')
@login_required
def subjects():
    """Subjects management page"""
    subjects_list = data_manager.get_subjects()
    return render_template('subjects.html', subjects=subjects_list)

@app.route('/subjects/add', methods=['POST'])
@login_required
def add_subject():
    """Add a new subject"""
    try:
        subject = {
            "name": request.form.get('name'),
            "type": request.form.get('type'),
            "code": request.form.get('code'),
            "hours_per_week": int(request.form.get('hours_per_week', 3)),
            "year": int(request.form.get('year', 1)),
            "semester": int(request.form.get('semester', 1))
        }
        data_manager.add_subject(subject)
        flash('Subject added successfully!', 'success')
    except Exception as e:
        flash(f'Error adding subject: {str(e)}', 'error')
    
    return redirect(url_for('subjects'))

@app.route('/subjects/edit/<int:subject_id>', methods=['POST'])
@login_required
def edit_subject(subject_id):
    """Edit/update an existing subject"""
    try:
        subject = {
            "name": request.form.get('name'),
            "type": request.form.get('type'),
            "code": request.form.get('code'),
            "hours_per_week": int(request.form.get('hours_per_week', 3)),
            "year": int(request.form.get('year', 1)),
            "semester": int(request.form.get('semester', 1))
        }
        data_manager.update_subject(subject_id, subject)
        flash('Subject updated successfully!', 'success')
    except Exception as e:
        flash(f'Error updating subject: {str(e)}', 'error')
    
    return redirect(url_for('subjects'))

@app.route('/subjects/delete/<int:subject_id>')
@login_required
def delete_subject(subject_id):
    """Delete a subject"""
    try:
        data_manager.delete_subject(subject_id)
        flash('Subject deleted successfully!', 'success')
    except Exception as e:
        flash(f'Error deleting subject: {str(e)}', 'error')
    
    return redirect(url_for('subjects'))

@app.route('/faculty')
@login_required
def faculty():
    """Faculty management page"""
    faculty_list = data_manager.get_faculty()
    subjects_list = data_manager.get_subjects()
    structure = data_manager.get_academic_structure()

    # Group subjects by academic year for simpler rendering in the template
    subjects_by_year = {}
    years_with_subjects = []
    for subject in subjects_list:
        try:
            year_value = subject.get('year') if isinstance(subject, dict) else getattr(subject, 'year', None)
        except Exception:
            year_value = None
        if year_value is None:
            continue
        if year_value not in subjects_by_year:
            subjects_by_year[year_value] = []
            years_with_subjects.append(year_value)
        subjects_by_year[year_value].append(subject)

    years_with_subjects = sorted(years_with_subjects)

    # Divisions per year from configuration
    divisions_per_year = int(structure.get('divisions_per_year', 2) or 2)
    available_divisions = list(range(1, divisions_per_year + 1))
    # Batches per division (for practicals) and user-friendly letters A, B, C, ...
    batches_per_division = int(structure.get('batches_per_division', 3) or 3)
    available_batches = list(range(1, batches_per_division + 1))
    available_batch_letters = [chr(ord('A') + i) for i in range(batches_per_division)]

    # Build a quick lookup for pre-filling edit modals: member_id -> {subject_id: divisions}
    faculty_assignments_map = {}
    for member in faculty_list:
        subject_to_divisions = {}
        assignments = member.get('subjects', []) or []
        for a in assignments:
            if isinstance(a, dict):
                sid = a.get('subject_id')
                divisions = a.get('divisions', []) or []
                batches = a.get('batches', []) or []
            else:
                sid = a
                divisions = []
                batches = []
            if sid is not None:
                subject_to_divisions[sid] = {"divisions": divisions, "batches": batches}
        faculty_assignments_map[member.get('id')] = subject_to_divisions

    # Compute hours/week for each faculty based on assigned subjects
    # Shows max of odd/even semester hours (faculty only teaches one semester at a time)
    # Theory/Tutorial: hours × division_count (same lecture repeated per division)
    # Practicals: hours × division_count (faculty teaches 1 batch per division, not all batches)
    subject_lookup_map = {s.get('id'): s for s in subjects_list}
    for member in faculty_list:
        odd_hours = 0
        even_hours = 0
        for a in member.get('subjects', []) or []:
            if isinstance(a, dict):
                sid = a.get('subject_id')
                subj = subject_lookup_map.get(sid, {})
                base_hours = int(subj.get('hours_per_week', 0) or 0)
                div_count = len(a.get('divisions', []))
                hrs = base_hours * max(div_count, 1)
                if subj.get('semester') == 1:
                    odd_hours += hrs
                else:
                    even_hours += hrs
            else:
                sid = a
                subj = subject_lookup_map.get(sid, {})
                hrs = int(subj.get('hours_per_week', 0) or 0)
                if subj.get('semester') == 1:
                    odd_hours += hrs
                else:
                    even_hours += hrs
        member['hours_per_week'] = max(odd_hours, even_hours)

    return render_template(
        'faculty.html',
        faculty=faculty_list,
        subjects=subjects_list,
        subjects_by_year=subjects_by_year,
        years_with_subjects=years_with_subjects,
        available_divisions=available_divisions,
        available_batches=available_batches,
        available_batch_letters=available_batch_letters,
        faculty_assignments_map=faculty_assignments_map,
    )

@app.route('/faculty/add', methods=['POST'])
@login_required
def add_faculty():
    """Add a new faculty member"""
    try:
        # Collect selected subjects and their division assignments
        selected_subject_ids = [int(s) for s in request.form.getlist('subjects') if s]
        subjects_list = data_manager.get_subjects()
        subject_lookup = {s.get('id'): s for s in subjects_list}

        assignments = []
        for sid in selected_subject_ids:
            divisions_key = f'divisions_{sid}'
            divisions = [int(d) for d in request.form.getlist(divisions_key)]
            # Optional batches per subject (for practicals)
            batches_key = f'batches_{sid}'
            batches = [int(b) for b in request.form.getlist(batches_key)]
            subj = subject_lookup.get(sid, {})
            assignments.append({
                'subject_id': sid,
                'year': int(subj.get('year', 0) or 0),
                'semester': int(subj.get('semester', 0) or 0),
                'divisions': divisions,
                'batches': batches
            })

        # Compute hours/week from selected subjects (max of odd/even semester)
        odd_hours = 0
        even_hours = 0
        for a in assignments:
            sid = a.get('subject_id')
            subj = subject_lookup.get(sid, {})
            base_hours = int(subj.get('hours_per_week', 0) or 0)
            div_count = len(a.get('divisions', []))
            hrs = base_hours * max(div_count, 1)
            if subj.get('semester') == 1:
                odd_hours += hrs
            else:
                even_hours += hrs
        hours_per_week = max(odd_hours, even_hours)

        faculty_member = {
            "name": request.form.get('name'),
            "initials": request.form.get('initials', '').upper(),
            "department": request.form.get('department'),
            "subjects": assignments,
            "hours_per_week": hours_per_week
        }

        data_manager.add_faculty(faculty_member)
        flash('Faculty member added successfully!', 'success')
    except Exception as e:
        flash(f'Error adding faculty member: {str(e)}', 'error')
    
    return redirect(url_for('faculty'))

@app.route('/faculty/edit/<int:faculty_id>', methods=['POST'])
@login_required
def edit_faculty(faculty_id):
    """Edit/update an existing faculty member"""
    try:
        selected_subject_ids = [int(s) for s in request.form.getlist('subjects') if s]
        subjects_list = data_manager.get_subjects()
        subject_lookup = {s.get('id'): s for s in subjects_list}

        assignments = []
        for sid in selected_subject_ids:
            divisions_key = f'divisions_{sid}'
            divisions = [int(d) for d in request.form.getlist(divisions_key)]
            batches_key = f'batches_{sid}'
            batches = [int(b) for b in request.form.getlist(batches_key)]
            subj = subject_lookup.get(sid, {})
            assignments.append({
                'subject_id': sid,
                'year': int(subj.get('year', 0) or 0),
                'semester': int(subj.get('semester', 0) or 0),
                'divisions': divisions,
                'batches': batches
            })

        # Compute hours/week from selected subjects (max of odd/even semester)
        odd_hours = 0
        even_hours = 0
        for a in assignments:
            sid = a.get('subject_id')
            subj = subject_lookup.get(sid, {})
            base_hours = int(subj.get('hours_per_week', 0) or 0)
            div_count = len(a.get('divisions', []))
            hrs = base_hours * max(div_count, 1)
            if subj.get('semester') == 1:
                odd_hours += hrs
            else:
                even_hours += hrs
        hours_per_week = max(odd_hours, even_hours)

        faculty_member = {
            "name": request.form.get('name'),
            "initials": request.form.get('initials', '').upper(),
            "department": request.form.get('department'),
            "subjects": assignments,
            "hours_per_week": hours_per_week
        }

        data_manager.update_faculty(faculty_id, faculty_member)
        flash('Faculty member updated successfully!', 'success')
    except Exception as e:
        flash(f'Error updating faculty member: {str(e)}', 'error')
    
    return redirect(url_for('faculty'))

@app.route('/faculty/delete/<int:faculty_id>')
@login_required
def delete_faculty(faculty_id):
    """Delete a faculty member"""
    try:
        data_manager.delete_faculty(faculty_id)
        flash('Faculty member deleted successfully!', 'success')
    except Exception as e:
        flash(f'Error deleting faculty member: {str(e)}', 'error')
    
    return redirect(url_for('faculty'))

@app.route('/rooms')
@login_required
def rooms():
    """Rooms management page"""
    rooms_list = data_manager.get_rooms()
    return render_template('rooms.html', rooms=rooms_list)

@app.route('/rooms/add', methods=['POST'])
@login_required
def add_room():
    """Add a new room"""
    try:
        room = {
            "name": request.form.get('name'),
            "type": request.form.get('type'),
            "capacity": int(request.form.get('capacity', 60)),
            "building": request.form.get('building', ''),
            "equipment": request.form.get('equipment', '')
        }
        data_manager.add_room(room)
        flash('Room added successfully!', 'success')
    except Exception as e:
        flash(f'Error adding room: {str(e)}', 'error')
    
    return redirect(url_for('rooms'))

@app.route('/rooms/delete/<int:room_id>')
@login_required
def delete_room(room_id):
    """Delete a room"""
    try:
        data_manager.delete_room(room_id)
        flash('Room deleted successfully!', 'success')
    except Exception as e:
        flash(f'Error deleting room: {str(e)}', 'error')
    
    return redirect(url_for('rooms'))

@app.route('/generate')
@login_required
def generate():
    """Timetable generation page"""
    return render_template('generate.html')

@app.route('/generate', methods=['POST'])
@login_required
def generate_timetable():
    """Generate timetable using AI solver"""
    try:
        # Read semester mode from form: 'odd' or 'even'
        semester_mode = request.form.get('semester_mode', 'odd')
        result = timetable_solver.generate_timetable(semester_mode=semester_mode)
        
        if "error" in result:
            flash(f'Error generating timetable: {result["error"]}', 'error')
            return redirect(url_for('generate'))

        # Show warnings if any
        if result.get("warnings"):
            for warning in result["warnings"]:
                flash(f'Warning: {warning}', 'warning')

        # Store the result for AI Assistant use
        if "timetable" in result:
            _timetable_cache['timetable'] = result["timetable"]
            session['current_timetable'] = result["timetable"]
            session['timetable_saved_at'] = datetime.now().isoformat()

            # Auto-save timetable (overwrites if same config fingerprint)
            try:
                data_manager.save_timetable(
                    result["timetable"],
                    semester_mode=semester_mode,
                    auto_save=True
                )
                flash('Timetable auto-saved.', 'info')
            except Exception as save_err:
                logging.error(f"Auto-save failed: {save_err}")

        return render_template('timetable.html', result=result, semester_mode=semester_mode)
        
    except Exception as e:
        logging.error(f"Error in generate_timetable: {str(e)}")
        flash(f'Error generating timetable: {str(e)}', 'error')
        return redirect(url_for('generate'))

@app.route('/api/timetable', methods=['POST'])
@login_required
def api_generate_timetable():
    """API endpoint for timetable generation"""
    try:
        result = timetable_solver.generate_timetable()
        return jsonify(result)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/current-timetable', methods=['GET'])
@login_required
def api_get_current_timetable():
    """API endpoint to get current timetable"""
    try:
        current_timetable = _get_current_timetable()
        if current_timetable:
            return jsonify({
                "success": True,
                "timetable": current_timetable,
                "saved_at": session.get('timetable_saved_at', '')
            })
        else:
            return jsonify({
                "success": False,
                "error": "No timetable available in session"
            })
    except Exception as e:
        logging.error(f"Error getting current timetable: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/structure')
@login_required
def api_get_structure():
    """API endpoint to get academic structure"""
    structure = data_manager.get_academic_structure()
    return jsonify(structure)

@app.route('/api/subjects')
@login_required
def api_get_subjects():
    """API endpoint to get subjects"""
    subjects_list = data_manager.get_subjects()
    return jsonify({"subjects": subjects_list})

@app.route('/api/faculty')
@login_required
def api_get_faculty():
    """API endpoint to get faculty"""
    faculty_list = data_manager.get_faculty()
    return jsonify({"faculty": faculty_list})

@app.route('/api/rooms')
@login_required
def api_get_rooms():
    """API endpoint to get rooms"""
    rooms_list = data_manager.get_rooms()
    return jsonify({"rooms": rooms_list})

# AI-Powered Timetable Modification Routes

@app.route('/ai-assistant')
@login_required
def ai_assistant_page():
    """AI Assistant page for timetable modifications"""
    return render_template('ai_assistant.html')

@app.route('/api/ai/process-request', methods=['POST'])
@login_required
def api_ai_process_request():
    """API endpoint for processing natural language timetable requests with feedback support"""
    try:
        data = request.get_json()
        user_request = data.get('request', '')
        user_feedback = data.get('user_feedback', '')
        current_timetable = _get_current_timetable()
        
        if not user_request:
            return jsonify({"error": "Request cannot be empty"}), 400
        
        # Get previous suggestions from session for context
        previous_suggestions = session.get('last_suggestions', [])
        
        # Process the request with feedback support
        result = ai_assistant.process_natural_language_request(
            user_request, 
            current_timetable, 
            user_feedback=user_feedback,
            previous_suggestions=previous_suggestions
        )
        
        # Store suggestions for future feedback
        if result.get('success') and result.get('reschedule_suggestions'):
            session['last_suggestions'] = result.get('reschedule_suggestions', {}).get('rescheduling_options', [])
        
        return jsonify(result)
        
    except Exception as e:
        logging.error(f"Error processing AI request: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/ai/analyze-conflicts', methods=['POST'])
@login_required
def api_ai_analyze_conflicts():
    """API endpoint for AI-powered conflict analysis"""
    try:
        data = request.get_json()
        timetable = data.get('timetable', {})
        
        if not timetable:
            return jsonify({"error": "Timetable data required"}), 400
        
        result = ai_assistant.analyze_schedule_conflicts(timetable)
        return jsonify(result)
        
    except Exception as e:
        logging.error(f"Error analyzing conflicts: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/ai/suggest-reschedule', methods=['POST'])
@login_required
def api_ai_suggest_reschedule():
    """API endpoint for AI-powered rescheduling suggestions"""
    try:
        data = request.get_json()
        affected_sessions = data.get('affected_sessions', [])
        constraints = data.get('constraints', {})
        
        result = ai_assistant.suggest_optimal_reschedule(affected_sessions, constraints)
        return jsonify(result)
        
    except Exception as e:
        logging.error(f"Error generating reschedule suggestions: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/ai/emergency-schedule', methods=['POST'])
@login_required
def api_ai_emergency_schedule():
    """API endpoint for emergency schedule generation"""
    try:
        data = request.get_json()
        constraints = data.get('constraints', {})
        priority_sessions = data.get('priority_sessions', [])
        
        result = ai_assistant.generate_emergency_schedule(constraints, priority_sessions)
        return jsonify(result)
        
    except Exception as e:
        logging.error(f"Error generating emergency schedule: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/ai/optimize-workload', methods=['POST'])
@login_required
def api_ai_optimize_workload():
    """API endpoint for faculty workload optimization"""
    try:
        data = request.get_json()
        current_assignments = data.get('assignments', {})
        
        result = ai_assistant.optimize_faculty_workload(current_assignments)
        return jsonify(result)
        
    except Exception as e:
        logging.error(f"Error optimizing workload: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/saved-timetables')
@login_required
def saved_timetables_page():
    """Saved timetables list page"""
    timetables = data_manager.get_saved_timetables()
    return render_template('saved_timetables.html', timetables=timetables)

@app.route('/saved-timetables/<int:timetable_id>')
@login_required
def saved_timetable_view(timetable_id):
    """View a specific saved timetable with Division/Batch/Faculty toggle"""
    saved = data_manager.get_saved_timetable(timetable_id)
    if not saved:
        flash('Timetable not found.', 'error')
        return redirect(url_for('saved_timetables_page'))
    return render_template('saved_timetables_view.html', saved=saved)

@app.route('/saved-timetables/delete/<int:timetable_id>', methods=['POST'])
@login_required
def delete_saved_timetable_page(timetable_id):
    """Delete a saved timetable and redirect back to list"""
    try:
        data_manager.delete_saved_timetable(timetable_id)
        flash('Timetable deleted successfully.', 'success')
    except Exception as e:
        flash(f'Error deleting timetable: {str(e)}', 'error')
    return redirect(url_for('saved_timetables_page'))

@app.route('/schedule-modifications')
@login_required
def schedule_modifications():
    """Schedule modifications management page"""
    return render_template('schedule_modifications.html')

@app.route('/api/save-timetable', methods=['POST'])
@login_required
def api_save_timetable():
    """Save current timetable to session for AI operations"""
    try:
        data = request.get_json()
        timetable = data.get('timetable', {})
        
        _timetable_cache['timetable'] = timetable
        session['current_timetable'] = timetable
        session['timetable_saved_at'] = datetime.now().isoformat()

        return jsonify({"success": True, "message": "Timetable saved successfully"})
        
    except Exception as e:
        logging.error(f"Error saving timetable: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/timetable/save-permanent', methods=['POST'])
@login_required
def api_save_timetable_permanent():
    """Save timetable permanently to database"""
    try:
        data = request.get_json()
        timetable = data.get('timetable', {})
        name = data.get('name', None)
        
        if not timetable or not timetable.keys():
            return jsonify({"error": "No timetable data provided"}), 400
        
        saved_entry = data_manager.save_timetable(timetable, name)
        
        return jsonify({
            "success": True,
            "message": "Timetable saved successfully",
            "timetable": saved_entry
        })
        
    except Exception as e:
        logging.error(f"Error saving timetable permanently: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/timetable/saved-list', methods=['GET'])
@login_required
def api_get_saved_timetables():
    """Get list of all saved timetables"""
    try:
        timetables = data_manager.get_saved_timetables()
        return jsonify({
            "success": True,
            "timetables": timetables
        })
    except Exception as e:
        logging.error(f"Error getting saved timetables: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/timetable/load-saved/<int:timetable_id>', methods=['POST'])
@login_required
def api_load_saved_timetable(timetable_id):
    """Load a saved timetable into session"""
    try:
        saved_timetable = data_manager.get_saved_timetable(timetable_id)
        
        if not saved_timetable:
            return jsonify({"error": "Timetable not found"}), 404
        
        # Load into cache and session
        _timetable_cache['timetable'] = saved_timetable['timetable']
        session['current_timetable'] = saved_timetable['timetable']
        session['timetable_saved_at'] = saved_timetable['saved_at']
        
        return jsonify({
            "success": True,
            "message": "Timetable loaded successfully",
            "timetable": saved_timetable['timetable'],
            "saved_at": saved_timetable['saved_at']
        })
        
    except Exception as e:
        logging.error(f"Error loading saved timetable: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/timetable/delete-saved/<int:timetable_id>', methods=['DELETE'])
@login_required
def api_delete_saved_timetable(timetable_id):
    """Delete a saved timetable"""
    try:
        data_manager.delete_saved_timetable(timetable_id)
        return jsonify({
            "success": True,
            "message": "Timetable deleted successfully"
        })
    except Exception as e:
        logging.error(f"Error deleting saved timetable: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/ai/get-suggestions', methods=['POST'])
@login_required
def api_ai_get_suggestions():
    """API endpoint to get 3 AI suggestions for a modification request"""
    try:
        data = request.get_json()
        user_request = data.get('request', '')
        current_timetable = _get_current_timetable()
        user_feedback = data.get('user_feedback', '')
        
        if not user_request:
            return jsonify({"error": "Request cannot be empty"}), 400
        
        # Process the request to get suggestions
        result = ai_assistant.get_modification_suggestions(
            user_request, 
            current_timetable, 
            user_feedback=user_feedback
        )
        
        return jsonify(result)
        
    except Exception as e:
        logging.error(f"Error getting suggestions: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/ai/apply-modification', methods=['POST'])
@login_required
def api_apply_modification():
    """Apply AI-suggested timetable modifications"""
    try:
        data = request.get_json()
        modification_type = data.get('modification_type')
        option = data.get('option', {})
        original_data = data.get('original_data', {})
        
        current_timetable = _get_current_timetable()
        if not current_timetable:
            return jsonify({"error": "No current timetable found. Please load a timetable first."}), 400
        
        # Apply the modification using AI assistant
        result = ai_assistant.apply_modification(modification_type, option, original_data, current_timetable)
        
        if result.get('success'):
            # Save the modified timetable back to cache and session
            if result.get('modified_timetable'):
                _timetable_cache['timetable'] = result['modified_timetable']
                session['current_timetable'] = result['modified_timetable']
                session['timetable_saved_at'] = datetime.now().isoformat()
                session['modification_history'] = session.get('modification_history', []) + [{
                    'timestamp': datetime.now().isoformat(),
                    'type': modification_type,
                    'description': option.get('description', 'Timetable modification'),
                    'applied_by': 'AI Assistant'
                }]
            
            # Mark as successful suggestion for tracking
            ai_assistant.mark_suggestion_successful()
        
        return jsonify(result)
        
    except Exception as e:
        logging.error(f"Error applying modification: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/ai/statistics', methods=['GET'])
@login_required
def get_ai_statistics():
    """Get AI assistant usage statistics and performance metrics"""
    try:
        stats = ai_assistant.get_ai_statistics()
        return jsonify({
            "success": True,
            "statistics": stats
        })
    except Exception as e:
        logging.error(f"Error getting AI statistics: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/ai/health', methods=['GET'])
@login_required
def get_ai_health():
    """Get AI assistant health status and capabilities"""
    try:
        health = ai_assistant.get_ai_health_status()
        return jsonify({
            "success": True,
            "health": health
        })
    except Exception as e:
        logging.error(f"Error getting AI health: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/ai/detect-conflicts', methods=['POST'])
@login_required
def detect_conflicts():
    """Use AI to detect schedule conflicts and suggest improvements"""
    try:
        data = request.get_json()
        timetable = data.get('timetable') or _get_current_timetable()
        
        if not timetable:
            return jsonify({"error": "No timetable data provided"}), 400
        
        result = ai_assistant.detect_schedule_conflicts_ai(timetable)
        return jsonify(result)
        
    except Exception as e:
        logging.error(f"Error detecting conflicts: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/timetable/export/pdf', methods=['POST'])
@login_required
def export_timetable_pdf():
    """Export ALL divisions/batches in timetable to PDF"""
    try:
        data = request.get_json()
        timetable = data.get('timetable')
        
        if not timetable:
            return jsonify({"error": "No timetable data provided"}), 400
        
        # Create PDF in memory
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), 
                              rightMargin=30, leftMargin=30, 
                              topMargin=30, bottomMargin=30)
        
        elements = []
        styles = getSampleStyleSheet()
        
        # Main Title
        title = Paragraph("<b>Complete Timetable - All Divisions</b>", styles['Title'])
        elements.append(title)
        elements.append(Spacer(1, 0.2*inch))
        
        # Add timestamp
        timestamp = Paragraph(f"Generated on: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}", 
                            styles['Normal'])
        elements.append(timestamp)
        elements.append(Spacer(1, 0.3*inch))
        
        # Sort views - divisions first (without "Batch"), then batches
        all_views = sorted(timetable.keys())
        division_views = [v for v in all_views if 'Batch' not in v]
        batch_views = [v for v in all_views if 'Batch' in v]
        sorted_views = division_views + batch_views
        
        # Create a page for each division/batch
        for idx, view_key in enumerate(sorted_views):
            view_data = timetable[view_key]
            view_name = view_key.replace('_', ' ')
            
            # Add page break between views (except first)
            if idx > 0:
                elements.append(PageBreak())
            
            # View title
            view_title = Paragraph(f"<b>{view_name}</b>", styles['Heading1'])
            elements.append(view_title)
            elements.append(Spacer(1, 0.2*inch))
            
            # Get data - viewData structure is {Day: {TimeSlot: [sessions]}}
            days = list(view_data.keys())
            day_order = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday']
            days.sort(key=lambda x: day_order.index(x) if x in day_order else 999)
            
            # Get time slots from first day
            time_slots = list(view_data[days[0]].keys()) if days and days[0] in view_data else []
            
            # Create table data
            table_data = []
            
            # Header row
            header = ['Time Slot'] + days
            table_data.append(header)
            
            # Data rows - skip continuation slots for practicals
            for time_slot in time_slots:
                row = [time_slot]
                for day in days:
                    # Get sessions for this day and time slot
                    slot_sessions = view_data.get(day, {}).get(time_slot, [])

                    # Ensure it's a list
                    if not isinstance(slot_sessions, list):
                        slot_sessions = [slot_sessions] if slot_sessions else []

                    if slot_sessions:
                        cell_text = []
                        for session in slot_sessions:
                            if session:  # Check session is not None
                                # Skip continuation slots - they're part of a multi-hour practical
                                if session.get('slot_position') == 'continuation':
                                    cell_text.append('↑ (contd.)')
                                    continue

                                session_type = session.get('type', '')

                                # Handle practical_block type (division view)
                                if session_type == 'practical_block':
                                    span = session.get('span', 1)
                                    duration_label = f" ({span} hrs)" if span > 1 else ""
                                    cell_text.append(f"Practicals{duration_label}")
                                    batches = session.get('batches', {})
                                    for b_num in sorted(batches.keys(), key=lambda x: int(x)):
                                        b_letter = chr(64 + int(b_num))
                                        b_info = batches[b_num]
                                        if b_info:
                                            cell_text.append(f"Batch {b_letter}: {b_info.get('subject', '')} ({b_info.get('room', '')})")
                                        else:
                                            cell_text.append(f"Batch {b_letter}: Free")
                                    cell_text.append('')
                                    continue

                                subject = session.get('subject', '')
                                faculty = session.get('faculty', '')
                                room = session.get('room', '')
                                span = session.get('span', 1)

                                duration_label = f" ({span} hrs)" if span > 1 else ""

                                if session_type:
                                    cell_text.append(f"{subject} ({session_type}){duration_label}")
                                else:
                                    cell_text.append(f"{subject}{duration_label}")

                                if faculty:
                                    cell_text.append(f"Faculty: {faculty}")
                                if room:
                                    cell_text.append(f"Room: {room}")
                                cell_text.append('')  # Empty line between sessions

                        row.append('\n'.join(cell_text).strip())
                    else:
                        row.append('-')

                table_data.append(row)
            
            # Create table
            table = Table(table_data, repeatRows=1)
            
            # Style the table
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#343a40')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('FONTSIZE', (0, 1), (-1, -1), 8),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
            ]))
            
            elements.append(table)
            elements.append(Spacer(1, 0.2*inch))
        
        # Build PDF
        doc.build(elements)
        buffer.seek(0)
        
        # Generate filename
        filename = f"Complete_Timetable_All_Divisions_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        
        return send_file(buffer, 
                        mimetype='application/pdf',
                        as_attachment=True,
                        download_name=filename)
        
    except Exception as e:
        logging.error(f"Error exporting PDF: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/timetable/export/excel', methods=['POST'])
@login_required
def export_timetable_excel():
    """Export ALL divisions/batches in timetable to Excel with separate sheets"""
    try:
        data = request.get_json()
        timetable = data.get('timetable')
        
        if not timetable:
            return jsonify({"error": "No timetable data provided"}), 400
        
        # Create workbook
        wb = Workbook()
        # Remove default sheet
        wb.remove(wb.active)
        
        # Define enhanced styles (similar to reference image)
        # Header style - Teal/Green color like reference image
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(start_color="2C7873", end_color="2C7873", fill_type="solid")
        
        # Time column style
        time_font = Font(bold=True, size=11, color="FFFFFF")
        time_fill = PatternFill(start_color="3D5A59", end_color="3D5A59", fill_type="solid")
        
        # Break style (for merged break cells)
        break_font = Font(bold=True, size=11, color="FFFFFF")
        break_fill = PatternFill(start_color="5A7D7A", end_color="5A7D7A", fill_type="solid")
        
        # Cell fills for theory/practical
        theory_fill = PatternFill(start_color="E8F4FD", end_color="E8F4FD", fill_type="solid")
        practical_fill = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
        free_fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")

        # Regular cell style
        regular_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left_alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        break_alignment = Alignment(horizontal='center', vertical='center')
        
        # Border styles
        border = Border(
            left=Side(style='thin', color='000000'),
            right=Side(style='thin', color='000000'),
            top=Side(style='thin', color='000000'),
            bottom=Side(style='thin', color='000000')
        )
        
        # Sort views - divisions first (without "Batch"), then batches
        all_views = sorted(timetable.keys())
        division_views = [v for v in all_views if 'Batch' not in v]
        batch_views = [v for v in all_views if 'Batch' in v]
        sorted_views = division_views + batch_views
        
        # Helper: get column letter for 1-based index (supports beyond Z)
        def col_letter(idx):
            """Convert 1-based column index to Excel column letter(s)."""
            from openpyxl.utils import get_column_letter
            return get_column_letter(idx)

        # Helper function to check if a time slot is a break
        def is_break_slot(time_slot, days_data, days_list):
            first_session = days_data.get(days_list[0], {}).get(time_slot)
            if not first_session:
                return None
            if isinstance(first_session, list):
                first_session = first_session[0] if first_session else None
            if not first_session or not isinstance(first_session, dict):
                return None
            if first_session.get('type') == 'break':
                return first_session.get('subject', 'Break')
            break_keywords = ['break', 'lunch', 'recess']
            first_subject = first_session.get('subject', '').lower()
            if any(kw in first_subject for kw in break_keywords):
                return first_session.get('subject', 'Break')
            return None

        # Thin border for inner practical cells
        thin_border = Border(
            left=Side(style='thin', color='999999'),
            right=Side(style='thin', color='999999'),
            top=Side(style='thin', color='999999'),
            bottom=Side(style='thin', color='999999')
        )
        faculty_font = Font(size=9, italic=True)

        # Create a sheet for each division/batch
        for view_key in sorted_views:
            view_data = timetable[view_key]
            view_name = view_key.replace('_', ' ')
            sheet_name = view_name[:31]
            ws = wb.create_sheet(title=sheet_name)

            days = list(view_data.keys())
            day_order = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday']
            days.sort(key=lambda x: day_order.index(x) if x in day_order else 999)
            time_slots = list(view_data[days[0]].keys()) if days and days[0] in view_data else []

            # Each day gets 2 sub-columns (subject | room)
            # Col 1 = Time, then for each day: left_col, right_col
            total_cols = 1 + len(days) * 2

            # Title row
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
            title_cell = ws.cell(row=1, column=1)
            title_cell.value = view_name
            title_cell.font = Font(bold=True, size=14, color="2C7873")
            title_cell.alignment = Alignment(horizontal='center', vertical='center')
            ws.row_dimensions[1].height = 25

            # Timestamp row
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
            ts_cell = ws.cell(row=2, column=1)
            ts_cell.value = f"Generated on: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}"
            ts_cell.font = Font(size=10, italic=True)
            ts_cell.alignment = Alignment(horizontal='center', vertical='center')
            ws.row_dimensions[2].height = 20

            # Header row (row 4) — Time + each day merged across 2 cols
            header_row = 4
            time_hdr = ws.cell(row=header_row, column=1)
            time_hdr.value = 'Time'
            time_hdr.font = header_font
            time_hdr.fill = header_fill
            time_hdr.alignment = Alignment(horizontal='center', vertical='center')
            time_hdr.border = border

            for d_idx, day in enumerate(days):
                left_col = 2 + d_idx * 2
                right_col = left_col + 1
                ws.merge_cells(start_row=header_row, start_column=left_col, end_row=header_row, end_column=right_col)
                hdr_cell = ws.cell(row=header_row, column=left_col)
                hdr_cell.value = day
                hdr_cell.font = header_font
                hdr_cell.fill = header_fill
                hdr_cell.alignment = Alignment(horizontal='center', vertical='center')
                hdr_cell.border = border
                ws.cell(row=header_row, column=right_col).border = border
            ws.row_dimensions[header_row].height = 25

            # Data rows
            current_row = header_row + 1

            for time_slot in time_slots:
                break_name = is_break_slot(time_slot, view_data, days)

                if break_name:
                    # Break row
                    time_cell = ws.cell(row=current_row, column=1)
                    time_cell.value = time_slot
                    time_cell.font = time_font
                    time_cell.fill = time_fill
                    time_cell.alignment = break_alignment
                    time_cell.border = border

                    ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=total_cols)
                    bc = ws.cell(row=current_row, column=2)
                    bc.value = break_name
                    bc.font = break_font
                    bc.fill = break_fill
                    bc.alignment = break_alignment
                    bc.border = border
                    for c in range(2, total_cols + 1):
                        ws.cell(row=current_row, column=c).border = border
                    ws.row_dimensions[current_row].height = 30
                    current_row += 1
                    continue

                # Determine how many sub-rows this time slot needs
                # Each practical batch needs 2 sub-rows (subject|room + faculty)
                # Theory needs 2 sub-rows (subject merged + faculty)
                max_batch_count = 0
                day_sessions = {}
                for day in days:
                    slot_data = view_data.get(day, {}).get(time_slot)
                    if isinstance(slot_data, list):
                        slot_data = slot_data[0] if slot_data else None
                    day_sessions[day] = slot_data
                    if slot_data and isinstance(slot_data, dict):
                        if slot_data.get('type') == 'practical_block':
                            batch_count = len(slot_data.get('batches', {}))
                            max_batch_count = max(max_batch_count, batch_count)

                # Each entry (theory or practical batch) takes 2 sub-rows
                # Use max(1, max_batch_count) entries, each 2 rows
                entries_count = max(1, max_batch_count)
                sub_rows = entries_count * 2

                # Time cell — merge vertically across all sub-rows
                ws.merge_cells(start_row=current_row, start_column=1,
                               end_row=current_row + sub_rows - 1, end_column=1)
                time_cell = ws.cell(row=current_row, column=1)
                time_cell.value = time_slot
                time_cell.font = time_font
                time_cell.fill = time_fill
                time_cell.alignment = Alignment(horizontal='center', vertical='center')
                time_cell.border = border
                for sr in range(sub_rows):
                    ws.cell(row=current_row + sr, column=1).border = border

                # Fill each day's columns
                for d_idx, day in enumerate(days):
                    left_col = 2 + d_idx * 2
                    right_col = left_col + 1
                    session = day_sessions[day]

                    if not session or not isinstance(session, dict) or session.get('type') == 'break':
                        # Free — merge both cols across all sub-rows
                        ws.merge_cells(start_row=current_row, start_column=left_col,
                                       end_row=current_row + sub_rows - 1, end_column=right_col)
                        fc = ws.cell(row=current_row, column=left_col)
                        fc.value = 'Free'
                        fc.fill = free_fill
                        fc.alignment = Alignment(horizontal='center', vertical='center')
                        fc.border = border
                        for sr in range(sub_rows):
                            ws.cell(row=current_row + sr, column=left_col).border = border
                            ws.cell(row=current_row + sr, column=right_col).border = border
                        continue

                    if session.get('slot_position') == 'continuation':
                        # Continuation — merge and show marker
                        ws.merge_cells(start_row=current_row, start_column=left_col,
                                       end_row=current_row + sub_rows - 1, end_column=right_col)
                        cc = ws.cell(row=current_row, column=left_col)
                        cc.value = '↑ (contd.)'
                        cc.fill = practical_fill
                        cc.alignment = Alignment(horizontal='center', vertical='center')
                        cc.border = border
                        for sr in range(sub_rows):
                            ws.cell(row=current_row + sr, column=left_col).border = border
                            ws.cell(row=current_row + sr, column=right_col).border = border
                        continue

                    session_type = session.get('type', '')

                    if session_type == 'practical_block':
                        batches = session.get('batches', {})
                        sorted_batches = sorted(batches.keys(), key=lambda x: int(x))
                        for b_idx, b_num in enumerate(sorted_batches):
                            b_info = batches[b_num]
                            row_offset = b_idx * 2
                            r1 = current_row + row_offset      # subject | room
                            r2 = current_row + row_offset + 1  # faculty (merged)

                            if b_info:
                                b_letter = chr(64 + int(b_num))
                                subj_name = b_info.get('subject', '')
                                room_name = b_info.get('room', '')
                                fac_name = b_info.get('faculty', '')

                                # Row 1: Subject (left) | Room (right)
                                sc = ws.cell(row=r1, column=left_col)
                                sc.value = f"{subj_name}"
                                sc.fill = practical_fill
                                sc.alignment = Alignment(horizontal='center', vertical='center')
                                sc.border = thin_border

                                rc = ws.cell(row=r1, column=right_col)
                                rc.value = room_name
                                rc.fill = practical_fill
                                rc.alignment = Alignment(horizontal='center', vertical='center')
                                rc.border = thin_border

                                # Row 2: Faculty (merged across both cols)
                                ws.merge_cells(start_row=r2, start_column=left_col,
                                               end_row=r2, end_column=right_col)
                                fcc = ws.cell(row=r2, column=left_col)
                                fcc.value = fac_name
                                fcc.font = faculty_font
                                fcc.fill = practical_fill
                                fcc.alignment = Alignment(horizontal='center', vertical='center')
                                fcc.border = thin_border
                                ws.cell(row=r2, column=right_col).border = thin_border
                            else:
                                # Free batch
                                ws.merge_cells(start_row=r1, start_column=left_col,
                                               end_row=r2, end_column=right_col)
                                fc = ws.cell(row=r1, column=left_col)
                                fc.value = 'Free'
                                fc.fill = free_fill
                                fc.alignment = Alignment(horizontal='center', vertical='center')
                                fc.border = thin_border

                        # Fill remaining sub-rows if this day has fewer batches
                        filled = len(sorted_batches)
                        for extra in range(filled, entries_count):
                            row_offset = extra * 2
                            r1 = current_row + row_offset
                            r2 = current_row + row_offset + 1
                            ws.merge_cells(start_row=r1, start_column=left_col,
                                           end_row=r2, end_column=right_col)
                            ec = ws.cell(row=r1, column=left_col)
                            ec.fill = free_fill
                            ec.border = thin_border

                    else:
                        # Theory/tutorial/individual practical — merge both cols AND all sub-rows
                        subject = session.get('subject', '')
                        faculty_name = session.get('faculty', '')
                        room_name = session.get('room', '')
                        span = session.get('span', 1)
                        duration_label = f" ({span} hrs)" if span > 1 else ""

                        fill = theory_fill if session_type in ('theory', 'tutorial') else practical_fill

                        # Merge entire area (both cols, all sub-rows) into one cell
                        r_end = current_row + sub_rows - 1
                        ws.merge_cells(start_row=current_row, start_column=left_col,
                                       end_row=r_end, end_column=right_col)
                        tc = ws.cell(row=current_row, column=left_col)
                        cell_lines = [f"{subject}{duration_label}"]
                        if faculty_name:
                            cell_lines.append(f"{faculty_name}")
                        if room_name:
                            cell_lines.append(f"{room_name}")
                        tc.value = '\n'.join(cell_lines)
                        tc.fill = fill
                        tc.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                        tc.border = thin_border
                        # Apply border to all cells in the merged range
                        for sr in range(sub_rows):
                            ws.cell(row=current_row + sr, column=left_col).border = thin_border
                            ws.cell(row=current_row + sr, column=right_col).border = thin_border

                # Set row heights for sub-rows
                for sr in range(sub_rows):
                    ws.row_dimensions[current_row + sr].height = 20

                current_row += sub_rows

            # Set column widths
            ws.column_dimensions[col_letter(1)].width = 14  # Time
            for d_idx in range(len(days)):
                left_c = 2 + d_idx * 2
                right_c = left_c + 1
                ws.column_dimensions[col_letter(left_c)].width = 18  # Subject
                ws.column_dimensions[col_letter(right_c)].width = 14  # Room
        
        # ---- Room View sheets ----
        # Build room schedules from timetable data
        room_schedules = {}  # { room_name: { day: { slot: { subject, faculty, type, context, batches } } } }
        for view_key, view_data in timetable.items():
            is_batch_view = 'Batch' in view_key
            context = view_key.replace('_', ' ')
            for day in view_data:
                slots = view_data[day]
                for time_slot, session in slots.items():
                    if not session or not isinstance(session, dict):
                        continue
                    if session.get('type') == 'break':
                        continue
                    if session.get('type') == 'practical_block':
                        if not is_batch_view:
                            continue
                        for b_num, b_info in (session.get('batches') or {}).items():
                            if not b_info or not b_info.get('room'):
                                continue
                            rname = b_info['room']
                            b_letter = chr(64 + int(b_num))
                            room_schedules.setdefault(rname, {}).setdefault(day, {})
                            if time_slot not in room_schedules[rname][day]:
                                room_schedules[rname][day][time_slot] = {
                                    'subject': b_info.get('subject', ''),
                                    'faculty': b_info.get('faculty', ''),
                                    'type': 'practical',
                                    'context': context,
                                    'batches': ['Batch ' + b_letter]
                                }
                            else:
                                existing = room_schedules[rname][day][time_slot]
                                existing.setdefault('batches', []).append('Batch ' + b_letter)
                                if b_info.get('faculty') and b_info['faculty'] not in existing.get('faculty', ''):
                                    existing['faculty'] = existing.get('faculty', '') + ', ' + b_info['faculty']
                    elif not is_batch_view:
                        rname = session.get('room')
                        if not rname:
                            continue
                        room_schedules.setdefault(rname, {}).setdefault(day, {})
                        if time_slot not in room_schedules[rname][day]:
                            room_schedules[rname][day][time_slot] = {
                                'subject': session.get('subject', ''),
                                'faculty': session.get('faculty', ''),
                                'type': session.get('type', 'theory'),
                                'context': context
                            }
                    else:
                        if session.get('type') in ('theory', 'tutorial'):
                            continue
                        rname = session.get('room')
                        if not rname:
                            continue
                        room_schedules.setdefault(rname, {}).setdefault(day, {})
                        if time_slot not in room_schedules[rname][day]:
                            room_schedules[rname][day][time_slot] = {
                                'subject': session.get('subject', ''),
                                'faculty': session.get('faculty', ''),
                                'type': session.get('type', 'practical'),
                                'context': context
                            }

        # Get standard day order and time slots
        sample_key = sorted_views[0] if sorted_views else None
        if sample_key:
            sample_view = timetable[sample_key]
            room_days = list(sample_view.keys())
            day_order = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday']
            room_days.sort(key=lambda x: day_order.index(x) if x in day_order else 999)
            room_time_slots = list(sample_view[room_days[0]].keys()) if room_days else []

            room_fill = PatternFill(start_color="E0F7FA", end_color="E0F7FA", fill_type="solid")

            for room_name in sorted(room_schedules.keys()):
                schedule = room_schedules[room_name]
                sheet_name = f"Room {room_name}"[:31]
                ws = wb.create_sheet(title=sheet_name)

                # Title
                merge_range = f'A1:{chr(65 + len(room_days))}1'
                ws.merge_cells(merge_range)
                title_cell = ws['A1']
                title_cell.value = f"Room: {room_name}"
                title_cell.font = Font(bold=True, size=14, color="2C7873")
                title_cell.alignment = Alignment(horizontal='center', vertical='center')
                ws.row_dimensions[1].height = 25

                # Timestamp
                merge_range_2 = f'A2:{chr(65 + len(room_days))}2'
                ws.merge_cells(merge_range_2)
                ts_cell = ws['A2']
                ts_cell.value = f"Generated on: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}"
                ts_cell.font = Font(size=10, italic=True)
                ts_cell.alignment = Alignment(horizontal='center', vertical='center')
                ws.row_dimensions[2].height = 20

                # Header
                header_row = 4
                ws.cell(row=header_row, column=1).value = 'Time'
                for col_idx, day in enumerate(room_days, start=2):
                    ws.cell(row=header_row, column=col_idx).value = day
                for col in range(1, len(room_days) + 2):
                    cell = ws.cell(row=header_row, column=col)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    cell.border = border
                ws.row_dimensions[header_row].height = 25

                # Data rows
                current_row = header_row + 1
                for time_slot in room_time_slots:
                    # Check break
                    first_session = timetable[sample_key].get(room_days[0], {}).get(time_slot)
                    is_break = first_session and isinstance(first_session, dict) and first_session.get('type') == 'break'

                    time_cell = ws.cell(row=current_row, column=1)
                    time_cell.value = time_slot
                    time_cell.font = time_font
                    time_cell.fill = time_fill
                    time_cell.alignment = Alignment(horizontal='center', vertical='center')
                    time_cell.border = border

                    if is_break:
                        merge_r = f'{chr(66)}{current_row}:{chr(65 + len(room_days))}{current_row}'
                        ws.merge_cells(merge_r)
                        bc = ws.cell(row=current_row, column=2)
                        bc.value = first_session.get('subject', 'Break')
                        bc.font = break_font
                        bc.fill = break_fill
                        bc.alignment = break_alignment
                        bc.border = border
                        for col in range(2, len(room_days) + 2):
                            ws.cell(row=current_row, column=col).border = border
                        ws.row_dimensions[current_row].height = 30
                    else:
                        max_lines = 1
                        for col_idx, day in enumerate(room_days, start=2):
                            cell = ws.cell(row=current_row, column=col_idx)
                            entry = schedule.get(day, {}).get(time_slot)
                            if entry:
                                lines = []
                                lines.append(f"{entry.get('subject', '')}")
                                lines.append(f"Faculty: {entry.get('faculty', '')}")
                                if entry.get('batches'):
                                    lines.append(f"{', '.join(entry['batches'])}")
                                lines.append(f"{entry.get('type', '').upper()} | {entry.get('context', '')}")
                                cell.value = '\n'.join(lines)
                                stype = entry.get('type', '')
                                if stype == 'practical':
                                    cell.fill = practical_fill
                                elif stype in ('theory', 'tutorial'):
                                    cell.fill = theory_fill
                                else:
                                    cell.fill = room_fill
                                max_lines = max(max_lines, len(lines))
                            else:
                                cell.value = 'Free'
                                cell.fill = free_fill
                            cell.alignment = left_alignment
                            cell.border = border
                        ws.row_dimensions[current_row].height = max(40, max_lines * 16)

                    current_row += 1

                # Column widths
                ws.column_dimensions['A'].width = 16
                for col_idx in range(2, len(room_days) + 2):
                    col_letter = chr(64 + col_idx)
                    ws.column_dimensions[col_letter].width = 35

        # Save to buffer
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        # Generate filename
        filename = f"Complete_Timetable_All_Divisions_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        return send_file(buffer,
                        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                        as_attachment=True,
                        download_name=filename)
        
    except Exception as e:
        logging.error(f"Error exporting Excel: {str(e)}")
        return jsonify({"error": str(e)}), 500
